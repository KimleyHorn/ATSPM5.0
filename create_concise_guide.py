from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path
import shutil

out=Path('ATSPM5_Account_Access_and_Administration_Guide.docx'); img=Path('atspm_guide_images'); img.mkdir(exist_ok=True)
new=[('00_signin_signup.png',r'C:\Users\JOSEPH~1.YOU\AppData\Local\Temp\codex-clipboard-af026fef-3200-4f57-a0ba-4fbeb05928c5.png'),('09_registration.png',r'C:\Users\joseph.young\OneDrive - KH\Pictures\Screenshots\Screenshot 2026-09-30 144625.png')]
for n,s in new: shutil.copy2(s,img/n)
doc=Document(); sec=doc.sections[0]; sec.top_margin=Inches(.6); sec.bottom_margin=Inches(.6); sec.left_margin=Inches(.7); sec.right_margin=Inches(.7)
for sn,sz in [('Normal',10),('Title',24),('Heading 1',17),('Heading 2',12)]:
 st=doc.styles[sn]; st.font.name='Aptos'; st.font.size=Pt(sz); st.font.color.rgb=RGBColor(0,0,0); st.font.bold=sn!='Normal'
doc.styles['Normal'].paragraph_format.space_after=Pt(5)
def shade(c,f):
 sh=OxmlElement('w:shd'); sh.set(qn('w:fill'),f); c._tc.get_or_add_tcPr().append(sh)
def tab(data):
 t=doc.add_table(rows=0,cols=len(data[0])); t.alignment=WD_TABLE_ALIGNMENT.CENTER
 for ri,row in enumerate(data):
  cs=t.add_row().cells
  for ci,v in enumerate(row):
   cs[ci].text=v; cs[ci].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   for p in cs[ci].paragraphs: p.paragraph_format.space_after=Pt(2); p.paragraph_format.space_before=Pt(2)
   if ri==0: shade(cs[ci],'D9E7F2')
   elif ri%2==0: shade(cs[ci],'F7F9FB')
 return t
def fig(fn,cap,w=6.4):
 p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(img/fn),width=Inches(w)); q=doc.add_paragraph(cap); q.alignment=WD_ALIGN_PARAGRAPH.CENTER; q.runs[0].italic=True; q.runs[0].font.size=Pt(8)
def bullets(items):
 for x in items: doc.add_paragraph('• '+x)

doc.add_paragraph('ATSPM 5 User Access Guide',style='Title')
doc.add_paragraph('A concise guide to signing up, requesting access, and administering roles',style='Subtitle')
doc.add_paragraph('This guide covers the normal user sign-up process and the administrator steps that follow. New users create their own accounts. Administrators do not need to create the initial account, but they may need to assign roles, geographic scope, or additional permissions after registration.')
doc.add_heading('1. New user sign-up',1)
doc.add_paragraph('From the ATSPM 5 sign-in screen, select Don’t have an account? Sign Up.')
fig('00_signin_signup.png','Figure 1. Select Don’t have an account? Sign Up from the sign-in screen.')
doc.add_paragraph('Complete the Registration form with the required information: First Name, Last Name, Email Address, Password, and Agency. Select Sign Up when the information is complete.')
fig('09_registration.png','Figure 2. Registration form for users creating their own account.')
bullets(['Use an agency email address whenever required by your organization.','Choose the correct agency; this may affect the administrator or scope that can approve access.','Save the sign-in email address and password securely.','If the account is not immediately able to view data or menus, contact the agency or system administrator for access assignment.'])
doc.add_heading('2. What happens after sign-up',1)
doc.add_paragraph('Registration creates the user account. It does not necessarily grant every ATSPM 5 function or every location. Access is controlled by roles and, where configured, by Regions, Jurisdictions, and Areas. An administrator should review the new account and assign only the access required for the user’s work.')
doc.add_heading('3. Administrator: assign access',1)
doc.add_paragraph('Open the administrator menu and select Users. Locate the user, open the Actions menu, and choose Edit. In User Details, verify the user’s identity, username, email, agency, and requested access.')
fig('03_user_details_form.png','Figure 3. User Details fields used to review and assign access.')
doc.add_paragraph('Use Roles to select one or more built-in or custom roles. Use Regions, Jurisdictions, and Areas to limit where the user can work or what locations the user can see. Select Save, then verify the user row in Manage Users.')
fig('02_user_details_roles.png','Figure 4. Role selector in User Details.')
doc.add_paragraph('The user list provides a quick review of Full Name, Username, Email, Agency, Roles, Regions, Jurisdictions, and Areas. The Actions menu provides Edit and Delete.')
fig('01_manage_users_updated.png','Figure 5. Manage Users list after an account is updated.')
doc.add_heading('4. Role reference',1)
tab([['Role','Use it for'],['Admin','Full access to configurations, data, reports, and user management.'],['User Admin','Managing user accounts and assigning roles.'],['Role Admin','Managing roles and permissions.'],['Data Admin','Accessing and exporting raw data logs.'],['Device Admin','Managing devices, products, and device configurations.'],['General Configuration Admin','Managing system-wide configuration except location-specific settings.'],['Location Configuration Admin','Managing location-specific settings and configurations.'],['Report Admin','Accessing restricted reports.'],['Usage Admin','Viewing, editing, and deleting usage reports.'],['Watchdog Subscriber','Viewing Watchdog reports and receiving Watchdog email notifications.'],['Api Key Admin / Speed Configuration Admin','Use only when the deployment confirms the specific need and permission.']])
doc.add_paragraph('Choose the least-privileged role that supports the user’s job. A role controls what the user may do; geographic scope controls where the user may do it.')
doc.add_heading('5. Create a custom role',1)
doc.add_paragraph('Use a custom role when a built-in role is broader than necessary. Open the administrator menu, select Roles, and choose New Custom Role.')
fig('06_manage_roles.png','Figure 6. Manage Roles includes built-in roles and Custom Roles.')
fig('05_create_role.png','Figure 7. Create New Role permission selectors.')
bullets(['Give the role a function-based name, such as Report Reviewer or Location Data Exporter.','Leave permission families at None unless the user needs them.','Add only the required permission level for User, Role, Location Configuration, General Configuration, Data, Watchdog, or Report.','Save the role, verify it appears under Custom Roles, and assign it from User Details.','Test the role with a non-production or test account when possible.'])
doc.add_heading('6. Quick troubleshooting',1)
tab([['Problem','Check first'],['User cannot see Users or Roles','Confirm User Admin or Role Admin is assigned, then have the user sign in again.'],['User sees no locations','Review Regions, Jurisdictions, and Areas.'],['User can view but cannot export','Review Data permission.'],['User cannot see a restricted report','Review Report permission and geographic scope.'],['Watchdog notifications are missing','Confirm Watchdog Subscriber, email address, and notification behavior.']])
doc.add_heading('7. Recommended additional screenshots',1)
bullets(['Account confirmation or approval message after self-registration.','Forgot-password and password-reset screens.','A sample user with limited geographic scope.','The Data export page and a restricted report.','The saved custom-role edit screen and any role-delete confirmation.'])
doc.add_paragraph('Administrator reminder: review access periodically, remove stale accounts or permissions, and avoid assigning Admin when a narrower role is sufficient.')
for s in doc.sections:
 f=s.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=f.add_run('ATSPM 5 User Access Guide  |  '); r.font.size=Pt(8); fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); f._p.append(fld)
doc.core_properties.title='ATSPM 5 User Access Guide'; doc.core_properties.author='ATSPM 5 Documentation'; doc.save(out)
print(out.resolve())
