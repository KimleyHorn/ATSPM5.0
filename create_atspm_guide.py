from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path
import shutil

OUT = Path('ATSPM5_Account_Access_and_Administration_Guide.docx')
IMGDIR = Path('atspm_guide_images')
IMGDIR.mkdir(exist_ok=True)
sources = [
('01_manage_users_updated.png', r'C:\Users\joseph.young\OneDrive - KH\Pictures\Screenshots\Screenshot 2026-09-30 171009.png'),
('02_user_details_roles.png', r'C:\Users\joseph.young\OneDrive - KH\Pictures\Screenshots\Screenshot 2026-09-30 170944.png'),
('03_user_details_form.png', r'C:\Users\joseph.young\OneDrive - KH\Pictures\Screenshots\Screenshot 2026-09-30 170935.png'),
('04_user_actions.png', r'C:\Users\joseph.young\OneDrive - KH\Pictures\Screenshots\Screenshot 2026-09-30 170930.png'),
('05_create_role.png', r'C:\Users\joseph.young\OneDrive - KH\Pictures\Screenshots\Screenshot 2026-09-30 170903.png'),
('06_manage_roles.png', r'C:\Users\joseph.young\OneDrive - KH\Pictures\Screenshots\Screenshot 2026-09-30 170856.png'),
('07_manage_users.png', r'C:\Users\joseph.young\OneDrive - KH\Pictures\Screenshots\Screenshot 2026-09-30 170508.png'),
('08_navigation_menu.png', r'C:\Users\joseph.young\OneDrive - KH\Pictures\Screenshots\Screenshot 2026-09-30 170452.png'),
]
for name, src in sources:
    print('copying', name, src, Path(src).exists())
    shutil.copy2(src, IMGDIR/name)

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(.65); sec.bottom_margin = Inches(.65); sec.left_margin = Inches(.75); sec.right_margin = Inches(.75)

styles = doc.styles
styles['Normal'].font.name = 'Aptos'; styles['Normal'].font.size = Pt(10); styles['Normal'].font.color.rgb = RGBColor(0,0,0)
styles['Normal'].paragraph_format.space_after = Pt(6)
for style_name, size in [('Title', 25), ('Heading 1', 18), ('Heading 2', 13), ('Heading 3', 11)]:
    st = styles[style_name]; st.font.name = 'Aptos Display' if style_name=='Title' else 'Aptos'; st.font.size = Pt(size); st.font.bold = True; st.font.color.rgb = RGBColor(0,0,0)
styles['Title'].paragraph_format.space_after = Pt(4)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr(); shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)
def borders(table, color='D9D9D9'):
    tblPr = table._tbl.tblPr; b = tblPr.first_child_found_in('w:tblBorders')
    if b is None: b = OxmlElement('w:tblBorders'); tblPr.append(b)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e = b.find(qn('w:'+edge))
        if e is None: e = OxmlElement('w:'+edge); b.append(e)
        e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:space'),'0'); e.set(qn('w:color'),color)
def table(data, widths=None, header=True):
    t=doc.add_table(rows=0, cols=len(data[0])); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'; borders(t)
    for ri,row in enumerate(data):
        cells=t.add_row().cells
        for ci,val in enumerate(row):
            cells[ci].text=str(val); cells[ci].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for p in cells[ci].paragraphs: p.paragraph_format.space_after=Pt(3); p.paragraph_format.space_before=Pt(3)
            if ri==0 and header: shade(cells[ci],'D9E7F2')
            elif ri%2==0: shade(cells[ci],'F7F9FB')
        if widths:
            for ci,w in enumerate(widths): cells[ci].width=Inches(w)
    return t
def caption(text):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(text); r.italic=True; r.font.size=Pt(8.5); r.font.color.rgb=RGBColor(80,80,80)
def figure(fname, cap, width=6.4):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(str(IMGDIR/fname), width=Inches(width)); caption(cap)
def bullet(text, level=0):
    p=doc.add_paragraph(style='List Bullet' if level==0 else 'List Bullet 2'); p.add_run(text); return p
def numbered(text):
    p=doc.add_paragraph(style='List Number'); p.add_run(text); return p

doc.add_paragraph('ATSPM 5 Account Access and Administration Guide', style='Title')
p=doc.add_paragraph(); r=p.add_run('User accounts, access scope, built-in roles, custom roles, and administrator procedures'); r.bold=True; r.font.size=Pt(12); r.font.color.rgb=RGBColor(55,95,130)
doc.add_paragraph('Version 1.0 | Prepared October 2026')
doc.add_paragraph('Purpose. This guide explains how ATSPM 5 users obtain access, how administrators manage users and permissions, what the displayed roles are intended to control, and how to create a custom role. It is written for agency administrators, system administrators, and users who need to understand why a menu or report may be unavailable to them.')
doc.add_paragraph('Evidence and scope. The procedures and role descriptions in this guide are based on the supplied ATSPM 5 interface screenshots. The screenshots show an administrator navigation menu, Manage Users, User Details, Manage Roles, and Create New Role. They do not show the initial public registration, invitation email, password-reset, or first-login screens; those steps are therefore identified as items to confirm in the deployed environment.')

doc.add_heading('1. Quick start', level=1)
table([
['If you are a…','Start here','Your outcome'],
['New user','Confirm your agency administrator has created or authorized your account, then sign in using the organization’s normal authentication process.','You can access the menus and data permitted by your assigned role and scope.'],
['User administrator','Open the administrator menu and choose Users. Create or edit the user, assign roles, and set Regions, Jurisdictions, and Areas as applicable.','The user appears in Manage Users with a visible access profile.'],
['Role administrator','Open the administrator menu and choose Roles. Review built-in roles or create a custom role with least-privilege permissions.','The role can be assigned to users and reused consistently.'],
['Troubleshooter','Compare the user’s Roles and scope fields with the menu or report they are trying to use.','You can distinguish a missing role from a missing location scope or a configuration issue.'],
], [1.35,3.2,2.0])

doc.add_heading('2. How ATSPM 5 access is organized', level=1)
doc.add_paragraph('The screenshots show two layers of authorization. First, a user receives one or more roles, such as Admin, Report Admin, or Data Admin. Second, the user may be limited to selected Regions, Jurisdictions, or Areas. A user can therefore have the correct type of permission but still be unable to see a particular location if that location is outside the user’s assigned scope.')
doc.add_paragraph('The administrator menu also provides the application’s navigation model. The visible groups are Location Configuration, Location Categories, User Management, and Other. The menu includes Locations, Routes, Products, Device Configurations, Areas, Regions, Jurisdictions, Users, Roles, FAQs, Menu Items, and Measure Defaults.')
figure('08_navigation_menu.png','Figure 1. Administrator navigation menu and the areas governed by different administrative permissions.')

doc.add_heading('3. User account lifecycle', level=1)
doc.add_heading('3.1 Account creation and first access', level=2)
doc.add_paragraph('The supplied screenshots begin after an administrator is already signed in. They do not show a self-service registration page or an invitation workflow. In a deployed ATSPM 5 environment, use the organization’s approved account-provisioning process and confirm whether accounts are created by an administrator, synchronized from an identity provider, or invited by email. The administrator should record the user’s agency, business purpose, requested scope, and approving authority before granting access.')
numbered('Collect the request: full name, username or email, agency, requested functions, and the Regions, Jurisdictions, or Areas the user needs.')
numbered('Verify the request with the agency or system owner. Do not grant the Admin role simply because a user needs to view a report.')
numbered('Create or locate the user in Users, then assign the smallest set of roles that supports the user’s work.')
numbered('Set the user’s geographic scope when the deployment uses scoped access. Leave scope broad only when the job function and agency policy require it.')
numbered('Ask the user to sign in and verify the expected menu items, reports, and locations. Record any exception for follow-up.')
doc.add_heading('3.2 User fields', level=2)
figure('03_user_details_form.png','Figure 2. User Details form. The visible fields are First Name, Last Name, Username, Agency, Email, Roles, Regions, Jurisdictions, and Areas.')
table([
['Field','How to use it','Quality check'],
['First Name / Last Name','Use the person’s agency-recognized name.','Avoid shared names or generic accounts unless formally approved.'],
['Username','Use the organization’s expected sign-in identifier.','Confirm spelling and case expectations before saving.'],
['Agency','Associate the user with the correct agency.','Verify the agency matches the requested locations and approval.'],
['Email','Use a monitored organizational address where policy requires it.','Check that alerts and account communications can reach the user.'],
['Roles','Assign built-in or custom permissions.','Use least privilege and document the reason for each elevated role.'],
['Regions / Jurisdictions / Areas','Limit access to the relevant geographic or organizational scope.','Confirm the selected scope contains the locations the user must access.'],
], [1.45,3.05,2.05])
doc.add_heading('3.3 Save, edit, and delete', level=2)
doc.add_paragraph('The User Details dialog includes Cancel and Save. After saving, return to Manage Users and verify that the row shows the expected name, username, email, agency, roles, and scope. The row’s vertical-ellipsis Actions menu provides Edit and Delete. Deletion is a consequential action: confirm the account is no longer needed, consider disabling or removing access according to agency policy, and preserve any required audit record before deleting.')
figure('04_user_actions.png','Figure 3. User Actions menu with Edit and Delete.')
figure('01_manage_users_updated.png','Figure 4. Manage Users after an update. The table exposes identity, agency, roles, and geographic scope at a glance.')

doc.add_heading('4. Assigning access to a user', level=1)
doc.add_paragraph('To grant access, open a user’s User Details record, select Roles, choose the required role or roles, select scope values when applicable, and save. The role picker shown in the supplied screenshot includes the following built-in roles: Admin, Api Key Admin, Data Admin, Device Admin, General Configuration Admin, Location Configuration Admin, Report Admin, Role Admin, Speed Configuration Admin, Usage Admin, User Admin, and Watchdog Subscriber.')
figure('02_user_details_roles.png','Figure 5. Role picker in User Details. Multiple roles can be selected for a single user.')
doc.add_paragraph('Use combinations deliberately. For example, a report consumer may need Report Admin or a narrower custom report role, while a person maintaining locations may need Location Configuration Admin. If a role is not visible in a user’s menus after assignment, check both the role assignment and geographic scope, then ask the user to sign out and back in if the application caches authorization state.')
doc.add_heading('4.1 Access-granting checklist', level=2)
for x in ['Requestor and approving authority are recorded.','User identity, agency, and email are correct.','Only required roles are selected.','Regions, Jurisdictions, and Areas are limited to the approved scope.','The user can access the required page or report after saving.','The administrator records the date, approver, and reason for elevated access.']:
    bullet('☐ '+x)

doc.add_heading('5. Built-in roles and what they do', level=1)
doc.add_paragraph('The Manage Roles screen provides the clearest descriptions of the built-in roles. The wording below follows the visible descriptions and adds an operational interpretation to help administrators choose the appropriate role. The operational interpretation should be confirmed against local policy and the current deployment because permissions can change as the application evolves.')
table([
['Role','Displayed purpose','Typical responsibility'],
['Admin','Full access to all configurations, data, reports, and user management.','System owner or tightly controlled superuser.'],
['Api Key Admin','Role is available in the user role picker; its detailed description is not visible in the supplied Manage Roles screenshot.','Manage API-key access only if the deployment confirms this permission.'],
['Data Admin','Can access and export raw data logs from the export page.','Data export, analysis, and controlled transfer of raw event data.'],
['Device Admin','Can manage devices, products, and device configurations.','Device and product configuration maintenance.'],
['General Configuration Admin','Can manage all system-wide configurations except location-specific settings.','Global settings, FAQs, areas, regions, jurisdictions, or other non-location-specific configuration as enabled.'],
['Location Configuration Admin','Can manage location-specific settings and configurations.','Locations, routes, and location-level configuration.'],
['Report Admin','Privileges are granted to access restricted reports.','Access to reports that are intentionally restricted.'],
['Role Admin','Can manage roles and permissions.','Create, edit, review, and assign permission sets.'],
['Speed Configuration Admin','Role is available in the user role picker; its detailed description is not visible in the supplied Manage Roles screenshot.','Speed-related configuration only if confirmed in the deployment.'],
['Usage Admin','Can view, edit, and delete usage reports.','Usage-report administration and cleanup.'],
['User Admin','Can manage Users.','Create, edit, assign, and remove user access.'],
['Watchdog Subscriber','Can view Watchdog reports and is subscribed to email notifications for Watchdog alerts.','Monitor system health and receive alert notifications.'],
], [1.55,2.8,2.2])
doc.add_heading('5.1 Important distinction: role versus scope', level=2)
doc.add_paragraph('A role answers “what may this user do?” Scope answers “where may this user do it?” Both must be correct. A user with Report Admin but no access to the relevant Jurisdiction may still be unable to view the expected report. Conversely, assigning a broad scope does not create permission to use a function that the user’s role does not allow.')

doc.add_heading('6. Creating and maintaining custom roles', level=1)
doc.add_paragraph('Use custom roles when a built-in role is broader than the user’s job requires. The Manage Roles page separates built-in roles from Custom Roles and provides a New Custom Role button. Custom roles make access easier to review and repeat because the same permission bundle can be assigned to multiple users.')
figure('06_manage_roles.png','Figure 6. Manage Roles separates built-in roles from Custom Roles and provides New Custom Role.')
doc.add_heading('6.1 Create a custom role', level=2)
numbered('Open the administrator menu and choose Roles.')
numbered('Select New Custom Role.')
numbered('Enter a clear Role Name that describes the job function, not a person’s name. Examples: Read Only Report Reviewer or Location Data Exporter.')
numbered('For each permission family, choose the narrowest permission level offered by the deployment. The visible families are User, Role, Location Configuration, General Configuration, Data, Watchdog, and Report.')
numbered('Review the description of each permission family before selecting it. For example, Data is described as exporting raw event logs; Watchdog includes viewing logs and subscribing to daily email updates; Report is described as viewing the left turn gap report.')
numbered('Select Save, then locate the new role in the Custom Roles table and verify that it appears with the expected name.')
figure('05_create_role.png','Figure 7. Create New Role dialog with permission families and level selectors.')
doc.add_heading('6.2 Custom-role design rules', level=2)
for x in ['Name roles by function and access level.','Start with all permission selectors at None, then add only what is needed.','Do not combine user-management and role-management permissions unless the job requires both.','Separate read-oriented reporting from export or configuration permissions.','Use a small number of well-defined roles rather than one custom role per person.','Test the role with a non-production test account before broad assignment.','Review custom roles whenever the application adds new permission families or when responsibilities change.']:
    bullet(x)

doc.add_heading('7. Administrator operating procedures', level=1)
doc.add_heading('7.1 New user request', level=2)
table([['Step','Administrator action','Evidence to retain'],['1','Validate requester, agency, business need, and approving authority.','Request or ticket number.'],['2','Create or edit the user record.','User Details values.'],['3','Assign role(s) and geographic scope.','Role and scope selections.'],['4','Save and verify the Manage Users row.','Updated user row or audit record.'],['5','Ask the user to test access.','Confirmation of expected page/report access.']], [0.55,3.6,2.4])
doc.add_heading('7.2 Access change', level=2)
doc.add_paragraph('Treat an access change as a new authorization decision. Confirm whether the change is temporary or permanent, modify only the affected role or scope, save, and test the result. When reducing access, verify that the user no longer sees the restricted function or data. When increasing access, verify that the added capability is available without unintentionally exposing unrelated locations or administrative functions.')
doc.add_heading('7.3 Offboarding', level=2)
doc.add_paragraph('When a user leaves an agency or no longer needs access, follow the organization’s retention and offboarding policy. The screenshots show Delete as an available user action, but they do not show a disable or suspension control. Confirm whether the deployed environment supports deactivation, and prefer the organization’s approved reversible action when available. Delete only after required records and approvals are preserved.')

doc.add_heading('8. Troubleshooting access problems', level=1)
table([['Symptom','Likely checks','Corrective action'],['User cannot see Users or Roles','User Admin or Role Admin is missing; account may require re-login.','Verify role assignment, sign out/in, and confirm the user is in the intended tenant or agency.'],['User sees a page but no locations','Regions, Jurisdictions, or Areas may be empty or too narrow.','Edit the user scope and retest with one known location.'],['User can view but cannot export','Data permission may be missing.','Assign a narrowly scoped data/export role only with approval.'],['User cannot access restricted reports','Report Admin or a custom report permission may be missing.','Review the report role and scope; do not grant Admin by default.'],['Watchdog email is not received','Watchdog Subscriber may be missing or email may be unverified.','Confirm the role, email address, subscription behavior, and notification settings.'],['Role assignment appears unchanged','Changes may not be saved or may be cached.','Reopen User Details, verify selections, save again, then reauthenticate.']], [1.55,2.75,2.25])

doc.add_heading('9. Recommended additional screenshots', level=1)
doc.add_paragraph('The supplied screenshots cover administrator-side management well. The following captures would make the guide complete and would remove the remaining deployment-specific uncertainty:')
for x in ['Public sign-in page and any Create Account or Request Access link.','Invitation or account-approval email, with personal information redacted.','First-login flow, including password creation, MFA, or identity-provider redirect if applicable.','Forgot-password and account-recovery screens.','The full Users page after selecting New User, if that action exists separately from editing.','Region, Jurisdiction, and Area pickers populated with example values, showing whether selections are hierarchical.','A user with a read-only role viewing the navigation menu, compared with an administrator view.','A restricted report that demonstrates the effect of Report Admin.','The Data export page and the confirmation or download state for Data Admin.','Watchdog reports and notification-subscription settings for Watchdog Subscriber.','Custom-role edit screen after saving, including how to modify or delete a custom role.','Any confirmation dialog, audit log, or toast message shown after deleting a user or role.']:
    bullet(x)

doc.add_heading('10. Administrator review checklist', level=1)
for x in ['Review built-in and custom roles quarterly.','Remove stale users and narrow overly broad scope.','Use named individual accounts wherever possible.','Document every Admin, Role Admin, User Admin, Data Admin, and API-key assignment.','Test access after role changes and after major application releases.','Redact personal information before sharing screenshots in tickets or training materials.','Keep this guide aligned with the live deployment if labels, roles, or permissions change.']:
    bullet('☐ '+x)

doc.add_heading('Appendix A. Screen-to-task index', level=1)
table([['Screenshot','Use in this guide'],['Manage Users updated','Confirm saved user identity, agency, roles, and scope.'],['User Details role picker','Select one or more built-in roles.'],['User Details form','Edit identity, agency, email, roles, and scope.'],['User Actions menu','Edit or delete a user.'],['Create New Role','Build a custom role from permission families.'],['Manage Roles','Review built-in roles and custom roles.'],['Manage Users overview','View the user directory and current assignments.'],['Administrator navigation','Find Users, Roles, and configuration areas.']], [2.2,4.35])
doc.add_paragraph('End of guide.')

# Footer with page numbering field
for section in doc.sections:
    footer=section.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=footer.add_run('ATSPM 5 Account Access and Administration Guide  |  '); run.font.size=Pt(8); run.font.color.rgb=RGBColor(100,100,100)
    fld=OxmlElement('w:fldSimple'); fld.set(qn('w:instr'),'PAGE'); footer._p.append(fld)

doc.core_properties.title='ATSPM 5 Account Access and Administration Guide'
doc.core_properties.subject='User accounts, roles, permissions, and administration'
doc.core_properties.author='ATSPM 5 Documentation'
doc.save(OUT)
print(OUT.resolve())
