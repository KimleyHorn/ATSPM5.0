from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, Image, Table, TableStyle
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.units import inch
from docx import Document
from pathlib import Path

src=Path('ATSPM5_Account_Access_and_Administration_Guide.docx')
out=Path('ATSPM5_Account_Access_and_Administration_Guide_Client.pdf')
docx=Document(src)
styles=getSampleStyleSheet()
styles.add(ParagraphStyle(name='Title2', parent=styles['Title'], fontName='Helvetica-Bold', fontSize=22, leading=26, textColor=colors.black, spaceAfter=12))
styles.add(ParagraphStyle(name='H1x', parent=styles['Heading1'], fontName='Helvetica-Bold', fontSize=16, leading=20, textColor=colors.black, spaceBefore=12, spaceAfter=6))
styles.add(ParagraphStyle(name='H2x', parent=styles['Heading2'], fontName='Helvetica-Bold', fontSize=12, leading=15, textColor=colors.black, spaceBefore=8, spaceAfter=4))
styles.add(ParagraphStyle(name='Bodyx', parent=styles['BodyText'], fontName='Helvetica', fontSize=9.2, leading=12, spaceAfter=5))
styles.add(ParagraphStyle(name='Smallx', parent=styles['BodyText'], fontSize=8, leading=10, textColor=colors.HexColor('#555555'), alignment=1))
story=[]
for p in docx.paragraphs:
    txt=p.text.strip()
    if not txt: continue
    st='Bodyx'
    if p.style.name=='Title': st='Title2'
    elif p.style.name=='Heading 1': st='H1x'
    elif p.style.name=='Heading 2': st='H2x'
    elif p.style.name.startswith('List'): txt='• '+txt
    story.append(Paragraph(txt.replace('&','&amp;').replace('<','&lt;').replace('>','&gt;'), styles[st]))
    # Add the supplied screenshots after the relevant figure-caption paragraphs.
    if txt.startswith('Figure '):
        n=int(txt.split()[1].strip('.'))
        fn={1:'00_signin_signup.png',2:'09_registration.png',3:'03_user_details_form.png',4:'02_user_details_roles.png',5:'01_manage_users_updated.png',6:'06_manage_roles.png',7:'05_create_role.png'}.get(n)
        if fn:
            im=Image(str(Path('atspm_guide_images')/fn)); im._restrictSize(6.5*inch,4.1*inch); story.append(im); story.append(Spacer(1,6))
for ti,t in enumerate(docx.tables):
    data=[]
    for row in t.rows:
        data.append([Paragraph(c.text.replace('&','&amp;').replace('<','&lt;').replace('>','&gt;'), styles['Bodyx']) for c in row.cells])
    tbl=Table(data, repeatRows=1, hAlign='LEFT')
    tbl.setStyle(TableStyle([('BACKGROUND',(0,0),(-1,0),colors.HexColor('#D9E7F2')),('GRID',(0,0),(-1,-1),.35,colors.HexColor('#D9D9D9')),('VALIGN',(0,0),(-1,-1),'MIDDLE'),('LEFTPADDING',(0,0),(-1,-1),5),('RIGHTPADDING',(0,0),(-1,-1),5),('TOPPADDING',(0,0),(-1,-1),4),('BOTTOMPADDING',(0,0),(-1,-1),4)]))
    story.append(tbl); story.append(Spacer(1,8))

def footer(canvas, doc):
    canvas.saveState(); canvas.setFont('Helvetica',8); canvas.setFillColor(colors.HexColor('#666666')); canvas.drawCentredString(letter[0]/2, .35*inch, f'ATSPM 5 Account Access and Administration Guide  |  {doc.page}'); canvas.restoreState()
SimpleDocTemplate(str(out), pagesize=letter, rightMargin=.6*inch,leftMargin=.6*inch,topMargin=.55*inch,bottomMargin=.55*inch).build(story,onFirstPage=footer,onLaterPages=footer)
print(out.resolve())
