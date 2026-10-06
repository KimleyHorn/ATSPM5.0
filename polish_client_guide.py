from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from pathlib import Path

path=Path('ATSPM5_Account_Access_and_Administration_Guide.docx')
doc=Document(path)

# Remove the entire optional recommendations section, including its bullets.
paras=doc.paragraphs
start=None; end=None
for i,p in enumerate(paras):
    if p.text.strip()=='7. Recommended additional screenshots': start=i
    if start is not None and p.text.strip().startswith('Administrator reminder:'): end=i; break
if start is not None:
    # Remove heading and all intervening bullets; retain the final administrator reminder.
    for p in paras[start:end if end is not None else len(paras)]:
        p._element.getparent().remove(p._element)

# Refine typography and color palette for a client-facing manual.
styles=doc.styles
styles['Normal'].font.name='Aptos'; styles['Normal'].font.size=Pt(10); styles['Normal'].font.color.rgb=RGBColor(32,32,32)
styles['Normal'].paragraph_format.space_after=Pt(6)
styles['Title'].font.name='Aptos Display'; styles['Title'].font.size=Pt(26); styles['Title'].font.bold=True; styles['Title'].font.color.rgb=RGBColor(0,0,0)
styles['Title'].paragraph_format.space_after=Pt(5)
styles['Subtitle'].font.name='Aptos'; styles['Subtitle'].font.size=Pt(12); styles['Subtitle'].font.color.rgb=RGBColor(49,103,160)
for n in ['Heading 1','Heading 2']:
    styles[n].font.name='Aptos'; styles[n].font.bold=True; styles[n].font.color.rgb=RGBColor(0,0,0)
styles['Heading 1'].font.size=Pt(17); styles['Heading 1'].paragraph_format.space_before=Pt(14); styles['Heading 1'].paragraph_format.space_after=Pt(6)
styles['Heading 2'].font.size=Pt(12)

# Add light gray borders to all tables and improve cell padding.
for table in doc.tables:
    tblPr=table._tbl.tblPr
    borders=tblPr.first_child_found_in('w:tblBorders')
    if borders is None: borders=OxmlElement('w:tblBorders'); tblPr.append(borders)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e=borders.find(qn('w:'+edge))
        if e is None: e=OxmlElement('w:'+edge); borders.append(e)
        e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:space'),'0'); e.set(qn('w:color'),'D9D9D9')
    for ri,row in enumerate(table.rows):
        for cell in row.cells:
            tc=cell._tc; tcPr=tc.get_or_add_tcPr(); mar=tcPr.first_child_found_in('w:tcMar')
            if mar is None: mar=OxmlElement('w:tcMar'); tcPr.append(mar)
            for side in ('top','start','bottom','end'):
                node=mar.find(qn('w:'+side))
                if node is None: node=OxmlElement('w:'+side); mar.append(node)
                node.set(qn('w:w'),'90'); node.set(qn('w:type'),'dxa')
            if ri==0:
                shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),'D9E7F2'); tcPr.append(shd)

# Make the footer understated and consistent.
for section in doc.sections:
    f=section.footer.paragraphs[0]; f.alignment=WD_ALIGN_PARAGRAPH.CENTER
    for r in f.runs: r.font.name='Aptos'; r.font.size=Pt(8); r.font.color.rgb=RGBColor(100,100,100)

doc.core_properties.title='ATSPM 5 User Access Guide'
doc.core_properties.subject='Client guide for self-registration and access administration'
doc.core_properties.author='ATSPM 5 Documentation'
doc.save(path)
print(path.resolve())
